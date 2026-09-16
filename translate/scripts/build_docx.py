# -*- coding: utf-8 -*-
"""Build a formatted academic translation DOCX from translation part files.

Usage:
    python build_docx.py <parts_dir> <out_docx>

<parts_dir> contains 01.json, 02.json, ... (sorted order). Each file:
    {"title_cn": "...",              # optional, first file only
     "blocks": [
        {"type":"title","text":"中文标题"},
        {"type":"subtitle","text":"English Original Title"},
        {"type":"heading","level":1,"text":"摘要"},
        {"type":"para","text":"..."},                       # body paragraph
        {"type":"para","text":"...","noindent":true},       # e.g. references entries
        {"type":"caption","text":"图1 ..."},
        {"type":"image","file":"images/p03_00.png","caption":"图1 ..."},  # file relative to parts_dir's parent
        {"type":"table","header":true,"rows":[["a","b"],["c","d"]]}
     ]}

Formatting: A4 single column; Chinese SimSun 10.5pt; Latin Times New Roman 10.5pt;
body justified, 1.5 line spacing, first-line indent 2 chars; headings bold black.
"""
import json, os, sys, glob
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

EA_FONT = "宋体"
LATIN_FONT = "Times New Roman"

def set_fonts(style, size_pt=None, bold=None, color_black=True):
    style.font.name = LATIN_FONT
    if size_pt is not None:
        style.font.size = Pt(size_pt)
    if bold is not None:
        style.font.bold = bold
    if color_black:
        style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), EA_FONT)

def run_fonts(run):
    run.font.name = LATIN_FONT
    r = run._element
    rpr = r.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), EA_FONT)

def setup_document(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(3.17)
    sec.right_margin = Cm(3.17)

    normal = doc.styles["Normal"]
    set_fonts(normal, 10.5)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)

    for name, size in (("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 10.5), ("Heading 4", 10.5)):
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        set_fonts(st, size, bold=True)
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing = 1.5
        st.paragraph_format.keep_with_next = True

def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run_fonts(run)
    run.bold = True
    run.font.size = Pt(15)

def add_subtitle(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(text)
    run_fonts(run)
    run.font.size = Pt(10.5)

def add_heading(doc, text, level):
    lvl = min(max(int(level), 1), 4)
    h = doc.add_heading("", level=lvl)
    run = h.add_run(text)
    run_fonts(run)

def _is_formula_para(text):
    """Display-formula paragraphs: LaTeX-heavy, little or no prose."""
    body = text.replace("$", " ").strip()
    cjk = sum(1 for ch in body if "\u4e00" <= ch <= "\u9fff")
    if cjk >= 2:
        return False
    if body.startswith("\\") or body.startswith("(") and "\\" in body:
        return True
    if "\\" in body and len(body) < 400:
        return True
    if text.count("$") >= 2 and len(body) < 120:
        return True
    return False

def add_para(doc, text, noindent=False):
    p = doc.add_paragraph()
    if _is_formula_para(text):
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if not noindent and not _is_formula_para(text):
        p.paragraph_format.first_line_indent = Pt(21)
    run = p.add_run(text)
    run_fonts(run)

def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run_fonts(run)
    run.font.size = Pt(9)

def add_image(doc, path, caption):
    if not os.path.exists(path):
        add_caption(doc, "（缺失图片：%s）" % os.path.basename(path))
        if caption:
            add_caption(doc, caption)
        return
    pic = doc.add_picture(path)
    width = pic.width if pic.width else 0
    height = pic.height if pic.height else 0
    max_w, max_h = Cm(14.6), Cm(20.0)
    if width > max_w or height > max_h:
        scale = min(max_w / width, max_h / height)
        pic.width = int(width * scale)
        pic.height = int(height * scale)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    if caption:
        add_caption(doc, caption)

def add_table(doc, rows, header):
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncols)
    try:
        t.style = doc.styles["Table Grid"]
    except KeyError:
        pass
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    cell_size = Pt(9) if ncols <= 6 else Pt(8)
    total_w = Cm(14.6)
    col_w = int(total_w / ncols)
    for j in range(ncols):
        for i in range(len(rows)):
            t.cell(i, j).width = col_w
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = t.cell(i, j)
            txt = str(row[j]) if j < len(row) else ""
            p = cell.paragraphs[0]
            run = p.add_run(txt)
            run_fonts(run)
            run.font.size = cell_size
            if header and i == 0:
                run.bold = True
    doc.add_paragraph()

def build(parts_dir, out_path):
    doc = Document()
    setup_document(doc)
    files = sorted(glob.glob(os.path.join(parts_dir, "*.json")))
    if not files:
        raise SystemExit("no part json files in " + parts_dir)
    for fp in files:
        with open(fp, encoding="utf-8") as f:
            data = json.load(f)
        blocks = data.get("blocks", [])
        for b in blocks:
            bt = b.get("type")
            if bt == "title":
                add_title(doc, b["text"])
            elif bt == "subtitle":
                add_subtitle(doc, b["text"])
            elif bt == "heading":
                add_heading(doc, b["text"], b.get("level", 1))
            elif bt == "para":
                add_para(doc, b["text"], noindent=bool(b.get("noindent")))
            elif bt == "caption":
                add_caption(doc, b["text"])
            elif bt == "image":
                img = b.get("file", "")
                base = os.path.dirname(parts_dir)
                path = img if os.path.isabs(img) else os.path.join(base, img)
                add_image(doc, path, b.get("caption", ""))
            elif bt == "table":
                add_table(doc, b.get("rows", []), bool(b.get("header")))
    doc.save(out_path)
    print("saved:", out_path, file=sys.stderr)

if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
