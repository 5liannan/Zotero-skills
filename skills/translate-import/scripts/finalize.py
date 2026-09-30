# -*- coding: utf-8 -*-
"""终稿规范化 —— 让翻译流程「一次构建即终稿」。

定位
====
本模块**属于正常翻译流程**：`build_docx.py` 在 `doc.save()` 之前调用
`finalize_document()`，因此 `translate-import` 产出的 DOCX 一落地就是终稿——
公式是 Word 原生公式对象（构建时由 `omml.py` 转换）、带页眉页码、图注规范、
参考文献可点击跳转，且引用号与文献列表一一对应。

也可以独立跑在**外部来源**的 DOCX 上（此时等价于原 `docx-polish` 的排版部分）：

    python finalize.py 译文.docx                 # 原地覆盖（自动备份）
    python finalize.py 译文.docx out.docx        # 另存
    python finalize.py 译文.docx --check-only    # 只做参考文献一致性检查

规范项（全部幂等，可重复运行）
============================
1. 图注          `图 1：xxx` / `图1 xxx` → `图 1　xxx`（全角空格），9pt 居中；
                 反向修复：被误排成图注样式的正文引用句（「图 1 给出了…」）
                 恢复为正文样式（两端对齐 + 缩进 2 字符 + 正文行距）
2. 标题          H1 14pt / H2 12pt，黑色加粗
3. 页眉 / 页脚   页眉 = 文章短标题；页脚 = PAGE 域；均 9pt 居中
4. 交叉引用      文末 `[n] …` 条目加隐藏书签，正文 `[n]` / `[n, m]` 换内部超链接
5. 残留标记      形如「原文图 p07_02.png」的孤立段整段删除
6. 参考文献一致性 正文引用号 ↔ 文末条目编号一一对应；悬挂引用、未引用条目、
                 编号断号 → 记为 error（`strict=True` 时直接中断构建）

与 docx-polish 的分工
=====================
`docx-polish` 是本技能内置的**事后补救步骤**，用于「别处给的、公式已退化成纯文本 LaTeX」的旧
DOCX（LaTeX→OMML 的事后补转）。本模块**不做**这件事——正常流程里公式在构建时
就已经是原生公式对象，没有可补转的东西。
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import os
import re
import shutil
import sys
import time

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

# ============================================================ 终稿规则
# 这一组常量/正则在 docx-polish/scripts/optimize.py 里还有一份：
# 补救入口需可独立运行，
# 所以规则是有意各存一份的。一致性由 tests/test_rule_parity.py 守卫 ——
# 改这里就必须改那边，否则 CI 会红。
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"

# ---------------------------------------------------------------- 图注
CAP_RE = re.compile(r"^\s*图\s*(\d+)\s*[.．:：]?\s*(.*)$", re.S)
# 图注主体若以这些动词起笔，说明是正文引用句（「图 1 给出了…」），不是图注。
# 注意**不要**把「由 / 为 / 是 / 如 / 见」列进来——它们是图注的常见起笔
# （例如「由四个方位的激光雷达测量反演得到的温度廓线…」）。
CAP_VERB_RE = re.compile(
    r"^\s*(?:给出|显示|展示|表明|说明|可见|反映|描绘|绘制|示意|标注)")
CAP_MAX_LEN = 200

# ------------------------------------------- 图片源文件追踪标记（上游残留）
IMG_MARK_RE = re.compile(
    r"^\s*(?:原文图|原始图|原图|图片|图源)\s*[:：]?\s*\S+"
    r"\.(?:png|jpg|jpeg|bmp|gif|tif|tiff|emf|wmf)\s*$", re.I)
CAP_HEAD_RE = re.compile(r"^\s*图\s*\d")

# ------------------------------------------------------------ 参考文献
REF_HEAD_RE = re.compile(
    r"^\s*(参考文献|參考文獻|引用文献|References|REFERENCES)\s*[:：]?\s*$")
REF_ITEM_RE = re.compile(r"^\s*\[(\d+)\]\s")
CITE_FIND = re.compile(r"\[(\d+(?:\s*[,，、]\s*\d+)*)\]")
NUM_SPLIT = re.compile(r"\s*[,，、]\s*")
CITE_RANGE_RE = re.compile(r"\[\s*\d+\s*[-–—]\s*\d+\s*\]")

HEADING_SIZES = {"Heading 1": 14, "Heading 2": 12,
                 "Heading 3": 10.5, "Heading 4": 10.5}

# 图注排版（同样由 test_rule_parity.py 与 docx-polish/optimize.py 对齐）
CAP_MAX_LEN = 200                # 超过这个长度不可能是图注
CAPTION_SIZE = 9.0
CAPTION_LINE_SPACING = 1.3
CAPTION_SPACE_BEFORE = 3
CAPTION_SPACE_AFTER = 8

BODY_LINE_SPACING = 1.5          # 正文行距倍数
INDENT_CHARS = 2.0               # 正文首行缩进字符数

EA_FONT = "宋体"
LATIN_FONT = "Times New Roman"


# ================================================================ 基础工具
def full_text(p):
    """段落全文——**含超链接内的 run**。

    `Paragraph.text` 不含 `w:hyperlink` 里的 run，用它扫引用标记会
    「整段丢掉引用编号且不报错」。所有文本判读一律走这里。
    """
    return "".join(t.text or "" for t in p._p.iter(qn("w:t")))


def _clear_paragraph(p):
    """清空段落内容（保留 pPr 段落属性）。"""
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)


def _apply_run(run, size=None, bold=None, ea=EA_FONT, latin=LATIN_FONT):
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), ea)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    return run


def _has_object(p):
    """段落里已有公式对象或图片时不要重建它的 run。"""
    if p._p.findall(".//{%s}oMath" % M_NS):
        return True
    if p._p.findall(".//" + qn("w:drawing")):
        return True
    return False


# ================================================================ 单段规范
def normalize_caption(text):
    """把图注规范成「图 N　题注」。

    是图注 → 返回规范化文本；不是图注（正文引用句 / 非图注段）→ 返回 None。
    """
    t = (text or "").strip()
    if not t or len(t) > CAP_MAX_LEN:
        return None
    m = CAP_RE.match(t)
    if not m:
        return None
    body = m.group(2).strip().lstrip("：:．.").strip()
    if not body:
        return None
    if CAP_VERB_RE.match(body):
        return None                      # 正文引用句，不是图注
    return "图 %s\u3000%s" % (m.group(1), body)


BODY_SIZE = 10.5


def _is_caption_shaped_prose(text):
    """是否是「图 N 给出了…」这类**正文引用句**（形似图注、实为正文）。"""
    m = CAP_RE.match((text or "").strip())
    if not m:
        return False
    body = m.group(2).strip().lstrip("：:．.").strip()
    return bool(body) and bool(CAP_VERB_RE.match(body))


def _looks_like_caption_style(p, body_size=BODY_SIZE):
    """段落当前是否被排成了图注样式（居中，或字号明显小于正文）。"""
    if p.alignment == WD_ALIGN_PARAGRAPH.CENTER:
        return True
    for r in p.runs:
        if r.font.size and r.font.size.pt < body_size - 0.5:
            return True
    return False


def _apply_body_style(p, body_size=BODY_SIZE, line_spacing=BODY_LINE_SPACING,
                      ea=EA_FONT, latin=LATIN_FONT):
    """按正文学体式重排（两端对齐 + 首行缩进 2 字符 + 正文行距）。"""
    pf = p.paragraph_format
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent = Pt(body_size * INDENT_CHARS)
    pf.line_spacing = line_spacing
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    for r in p.runs:
        _apply_run(r, size=body_size, ea=ea, latin=latin)
    return p


def apply_captions(doc, body_size=BODY_SIZE, fix_prose=True,
                   ea=EA_FONT, latin=LATIN_FONT):
    """规范图注段落，并修掉「正文被误排成图注」的反向缺陷。

    正向：`图 1：xxx` / `图1 xxx` → 「图 1　xxx」，9pt 居中。

    反向：外部旧文档里常见上游流水线把**所有**「图 N」开头的段落都当图注排，
    于是「图 1 给出了…」这类正文引用句被排成居中 9pt。这类段落会被恢复成正文样式。

    保守起见，只改写**当前确实排成图注样**（居中或小字号）的段落——正常的正文
    段落不会因为开头恰好是「图 9 展示了…」而被动到。

    返回 {"captions": n, "restyled_prose": m}。
    """
    cap = prose = 0
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading"):
            continue
        if _has_object(p):
            continue
        t = full_text(p).strip()
        if not t:
            continue

        new = normalize_caption(t)
        if new:
            _clear_paragraph(p)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf = p.paragraph_format
            pf.first_line_indent = Pt(0)
            pf.line_spacing = CAPTION_LINE_SPACING
            pf.space_before = Pt(CAPTION_SPACE_BEFORE)
            pf.space_after = Pt(CAPTION_SPACE_AFTER)
            _apply_run(p.add_run(new), size=CAPTION_SIZE)
            cap += 1
            continue

        if (fix_prose and _is_caption_shaped_prose(t)
                and _looks_like_caption_style(p, body_size)):
            _apply_body_style(p, body_size, ea=ea, latin=latin)
            prose += 1

    return {"captions": cap, "restyled_prose": prose}


def apply_headings(doc, ea=EA_FONT, latin=LATIN_FONT):
    """统一标题层级字号（样式 + run），返回处理条数。"""
    for name, size in HEADING_SIZES.items():
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        st.font.name = latin
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), ea)

    n = 0
    for p in doc.paragraphs:
        if not p.style.name.startswith("Heading"):
            continue
        sz = HEADING_SIZES.get(p.style.name, 12)
        for r in p.runs:
            _apply_run(r, size=sz, bold=True, ea=ea, latin=latin)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = Pt(0)
        n += 1
    return n


def strip_image_marks(doc):
    """删除「原文图 p07_02.png」这类图片源文件追踪残留段，返回删除条数。"""
    n = 0
    for p in list(doc.paragraphs):
        t = full_text(p).strip()
        if not t or CAP_HEAD_RE.match(t):
            continue
        if IMG_MARK_RE.match(t):
            parent = p._element.getparent()
            if parent is not None:
                parent.remove(p._element)
                n += 1
    return n


# ================================================================ 页眉页脚
def set_header(doc, title, size=9, ea=EA_FONT, latin=LATIN_FONT):
    sec = doc.sections[0]
    hdr = sec.header
    try:
        hdr.is_linked_to_previous = False
    except Exception:
        pass
    p = hdr.paragraphs[0] if hdr.paragraphs else hdr.add_paragraph()
    _clear_paragraph(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if title:
        _apply_run(p.add_run(title), size=size, ea=ea, latin=latin)
    return bool(title)


def set_footer_page(doc, size=9, ea=EA_FONT, latin=LATIN_FONT):
    """页脚插入 PAGE 域。不能只写 'PAGE' 文本，必须构造 fldChar。"""
    sec = doc.sections[0]
    ftr = sec.footer
    try:
        ftr.is_linked_to_previous = False
    except Exception:
        pass
    p = ftr.paragraphs[0] if ftr.paragraphs else ftr.add_paragraph()
    _clear_paragraph(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = _apply_run(p.add_run(), size=size, ea=ea, latin=latin)
    r = run._element
    for kind, text in (("begin", None), (None, " PAGE "), ("end", None)):
        if kind:
            el = r.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): kind})
        else:
            el = r.makeelement(qn("w:instrText"), {qn("xml:space"): "preserve"})
            el.text = text
        r.append(el)
    return True


# ============================================================ 参考文献检查
def check_references(doc, allow_unused=False):
    """校验正文引用号 ↔ 文末文献条目编号是否一一对应。

    返回 dict：refs / cites / dangling / unused / gaps / errors / warnings。
    `errors` 非空即视为不合格（构建侧据此中断）。
    """
    ps = doc.paragraphs

    head = None
    for i, p in enumerate(ps):
        if REF_HEAD_RE.match(full_text(p).strip()):
            head = i
            break

    ref_map = {}
    if head is not None:
        started = False
        for i in range(head + 1, len(ps)):
            t = full_text(ps[i]).strip()
            m = REF_ITEM_RE.match(t)
            if m:
                ref_map[int(m.group(1))] = t
                started = True
            elif started and t:
                break            # 条目区结束（空段不算结束，避免夹缝空行误判）

    cites = {}
    ranges = []
    upto = head if head is not None else len(ps)
    for i in range(0, upto):
        t = full_text(ps[i])
        if not t:
            continue
        for m in CITE_FIND.finditer(t):
            for chunk in NUM_SPLIT.split(m.group(1)):
                if chunk.isdigit():
                    cites.setdefault(int(chunk), []).append(i)
        for m in CITE_RANGE_RE.finditer(t):
            ranges.append(m.group(0))

    refs = sorted(ref_map)
    cited = sorted(cites)
    errors, warnings = [], []

    if head is None:
        if cited:
            errors.append("正文出现引用标记 %s，但找不到「参考文献 / References」标题段" % cited)
    else:
        if not refs:
            errors.append("「参考文献」标题（段 %d）之后没有任何 `[n] …` 文献条目" % head)
        if refs and not cited:
            errors.append("文末有 %d 条文献条目，但正文一处引用标记都没有" % len(refs))

    dangling = [n for n in cited if n not in ref_map]
    if dangling:
        errors.append("悬挂引用（正文引用了、文献列表没有）：%s" % dangling)

    unused = [n for n in refs if n not in cites]
    if unused:
        msg = "未引用条目（文献列表有、正文未引用）：%s" % unused
        (warnings if allow_unused else errors).append(msg)

    gaps = [n for n in range(1, (max(refs) + 1) if refs else 1) if n not in ref_map]
    if gaps:
        errors.append("文献编号不连续，缺：%s" % gaps)

    if ranges:
        warnings.append("发现 %d 处区间式引用（如 [5-7]），交叉引用不会展开：%s"
                        % (len(ranges), sorted(set(ranges))[:6]))

    return {"head": head, "refs": len(refs), "cites": len(cited),
            "ref_nums": refs, "cite_nums": cited,
            "dangling": dangling, "unused": unused, "gaps": gaps,
            "range_notations": sorted(set(ranges)),
            "errors": errors, "warnings": warnings}


# ================================================================ 交叉引用
def _mk_run(rPr, text, link_style=False):
    r = OxmlElement("w:r")
    if rPr is not None:
        nrPr = copy.deepcopy(rPr)
        if link_style:
            for tag in ("w:color", "w:u"):
                el = nrPr.find(qn(tag))
                if el is not None:
                    nrPr.remove(el)
            nrPr.insert(0, OxmlElement("w:rStyle"))
            nrPr[0].set(qn("w:val"), "Hyperlink")
        r.append(nrPr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def _mk_link(rPr, text, anchor, link_style=False):
    h = OxmlElement("w:hyperlink")
    h.set(qn("w:anchor"), anchor)
    h.set(qn("w:history"), "1")
    h.append(_mk_run(rPr, text, link_style))
    return h


def _next_bookmark_id(doc):
    ids = []
    for e in doc.element.body.iter(qn("w:bookmarkStart")):
        v = e.get(qn("w:id"))
        if v and v.lstrip("-").isdigit():
            ids.append(int(v))
    return (max(ids) + 1) if ids else 1


def add_xref(doc, prefix="_RefList", link_style=False):
    """文献条目加隐藏书签 + 正文引用换内部超链接（编号文本不动）。

    为什么不用 `REF \\r \\h` 域：`\\r` 会把列表编号的格式一起带出来
    （编号格式 `[%1]` → 域结果 `[1]`），分组引用 `[5, 6]` 会变成 `[[5], [6]]`。
    纯超链接方案没有这个问题，代价是编号为死文本、增删文献不会自动重编号。
    """
    ps = doc.paragraphs
    head = None
    for i, p in enumerate(ps):
        if REF_HEAD_RE.match(full_text(p).strip()):
            head = i
            break
    if head is None:
        return {"refs": 0, "links": 0, "replaced": 0, "skipped": "no ref section"}

    ref_map = {}
    first = last = None
    for i in range(head + 1, len(ps)):
        m = REF_ITEM_RE.match(full_text(ps[i]).strip())
        if m:
            ref_map[int(m.group(1))] = ps[i]
            first = i if first is None else first
            last = i
        elif ref_map:
            break
    if not ref_map:
        return {"refs": 0, "links": 0, "replaced": 0, "skipped": "no ref items"}

    nums = sorted(ref_map)
    # 幂等：已有同名书签就不再重复添加（重复书签名是无效 OOXML，
    # Word 只会认第一个，还会让文档在校验/转换工具里报错）
    existing = {e.get(qn("w:name")) or ""
                for e in doc.element.body.iter(qn("w:bookmarkStart"))}
    bid = _next_bookmark_id(doc)
    added = 0
    for n in nums:
        name = "%s%d" % (prefix, n)
        if name in existing:
            continue
        p = ref_map[n]
        pPr = p._p.find(qn("w:pPr"))
        pos = (list(p._p).index(pPr) + 1) if pPr is not None else 0
        s = OxmlElement("w:bookmarkStart")
        s.set(qn("w:id"), str(bid))
        s.set(qn("w:name"), name)
        e = OxmlElement("w:bookmarkEnd")
        e.set(qn("w:id"), str(bid))
        p._p.insert(pos, s)
        p._p.append(e)
        bid += 1
        added += 1

    replaced = links = 0
    for i in range(0, head):
        for run in list(ps[i].runs):
            t = run.text
            if not t or "[" not in t:
                continue
            hits = list(CITE_FIND.finditer(t))
            if not hits:
                continue
            r = run._r
            parent = r.getparent()
            pos = list(parent).index(r)
            rPr = r.find(qn("w:rPr"))

            new = []
            cursor = 0
            ok = True
            for m in hits:
                groups = []
                for chunk in NUM_SPLIT.split(m.group(1)):
                    try:
                        groups.append(int(chunk))
                    except ValueError:
                        ok = False
                if not ok or any(x not in ref_map for x in groups):
                    ok = False
                    break
                if m.start() > cursor:
                    new.append(_mk_run(rPr, t[cursor:m.start()]))
                if len(groups) == 1:
                    new.append(_mk_link(rPr, "[%d]" % groups[0],
                                        "%s%d" % (prefix, groups[0]), link_style))
                    links += 1
                else:
                    new.append(_mk_run(rPr, "["))
                    for k, x in enumerate(groups):
                        if k:
                            new.append(_mk_run(rPr, ", "))
                        new.append(_mk_link(rPr, str(x), "%s%d" % (prefix, x),
                                            link_style))
                        links += 1
                    new.append(_mk_run(rPr, "]"))
                cursor = m.end()
            if not ok or not new:
                continue
            if cursor < len(t):
                new.append(_mk_run(rPr, t[cursor:]))
            for k, el in enumerate(new):
                parent.insert(pos + k, el)
            parent.remove(r)
            replaced += 1

    return {"refs": len(nums), "links": links, "replaced": replaced,
            "bookmarks_added": added, "bookmarks_total": len(nums),
            "skipped": None, "first": first, "last": last}


# ================================================================ 主入口
def finalize_document(doc, title=None, *, header=True, page_number=True,
                      captions=True, headings=True, strip_marks=True,
                      xref=True, check_refs=True, allow_unused_refs=False,
                      xref_style=False, strict=False, header_size=9,
                      ea=EA_FONT, latin=LATIN_FONT):
    """就地规范化一个已构建好的 Document。

    strict=True 时，参考文献一致性检查一旦有 error 立即抛 SystemExit
    （此时**尚未做任何修改**，调用方不会产出半成品文件）。
    """
    stats = {}

    ref = check_references(doc, allow_unused=allow_unused_refs) if check_refs else None
    stats["refs_check"] = ref
    if strict and ref and ref["errors"]:
        raise SystemExit(
            "参考文献一致性检查未通过，已中止写出：\n  - "
            + "\n  - ".join(ref["errors"]))

    if strip_marks:
        stats["marks_removed"] = strip_image_marks(doc)
    if captions:
        stats["captions"] = apply_captions(doc, ea=ea, latin=latin)
    if headings:
        stats["headings"] = apply_headings(doc, ea=ea, latin=latin)

    if title is None:
        for p in doc.paragraphs:
            t = full_text(p).strip()
            if t:
                title = t[:80]
                break

    if header:
        stats["header"] = set_header(doc, title or "", header_size, ea, latin)
    if page_number:
        stats["page_number"] = set_footer_page(doc, header_size, ea, latin)

    stats["xref"] = add_xref(doc, link_style=xref_style) if xref else None
    return stats


def _md5(path):
    h = hashlib.md5()
    with open(path, "rb") as f:
        for c in iter(lambda: f.read(1 << 20), b""):
            h.update(c)
    return h.hexdigest()


def backup_file(path):
    """带日期后缀的备份（同名则加时间戳），返回备份路径；失败返回 None。"""
    if not os.path.exists(path):
        return None
    bak = path + ".bak_" + time.strftime("%Y%m%d")
    if os.path.exists(bak):
        bak = path + ".bak_" + time.strftime("%Y%m%d_%H%M%S")
    shutil.copy2(path, bak)
    if os.path.getsize(bak) == os.path.getsize(path) and _md5(bak) == _md5(path):
        return bak
    return None


def finalize(path, out_path=None, *, backup=True, title=None, **kw):
    """对磁盘上的 DOCX 做终稿规范化。原地覆盖时先备份并双校验。"""
    doc = Document(path)
    stats = finalize_document(doc, title=title, **kw)
    dst = out_path or path
    if os.path.abspath(dst) == os.path.abspath(path):
        stats["backup"] = backup_file(path) if backup else None
    doc.save(dst)
    stats["out"] = dst
    return stats


# ===================================================================== CLI
def main():
    ap = argparse.ArgumentParser(
        description="译文 DOCX 终稿规范化（页眉页码 / 图注 / 标题 / 交叉引用 / 参考文献一致性）")
    ap.add_argument("src", help="输入 .docx")
    ap.add_argument("out", nargs="?", help="输出 .docx（默认原地覆盖，自动备份）")
    ap.add_argument("--title", default=None, help="页眉标题（默认取首个非空段落）")
    ap.add_argument("--no-header", action="store_true")
    ap.add_argument("--no-page-number", action="store_true")
    ap.add_argument("--no-captions", action="store_true")
    ap.add_argument("--no-headings", action="store_true")
    ap.add_argument("--no-strip-marks", action="store_true")
    ap.add_argument("--no-xref", action="store_true", help="不做交叉引用")
    ap.add_argument("--hyperlink-style", action="store_true",
                    help="交叉引用套用蓝色下划线样式（默认保持正文字体外观）")
    ap.add_argument("--allow-unused-refs", action="store_true",
                    help="文献列表中存在未被正文引用的条目时不算错误")
    ap.add_argument("--check-only", action="store_true", help="只做参考文献一致性检查，不改文件")
    ap.add_argument("--no-backup", action="store_true")
    ap.add_argument("--quiet", action="store_true")
    a = ap.parse_args()

    if a.check_only:
        ref = check_references(Document(a.src),
                               allow_unused=a.allow_unused_refs)
        print(json_dump(ref))
        sys.exit(1 if ref["errors"] else 0)

    stats = finalize(
        a.src, a.out,
        backup=not a.no_backup,
        title=a.title,
        header=not a.no_header,
        page_number=not a.no_page_number,
        captions=not a.no_captions,
        headings=not a.no_headings,
        strip_marks=not a.no_strip_marks,
        xref=not a.no_xref,
        xref_style=a.hyperlink_style,
        allow_unused_refs=a.allow_unused_refs,
    )
    if not a.quiet:
        print(json_dump(stats))
    if stats["refs_check"] and stats["refs_check"]["errors"]:
        print("注意：参考文献一致性检查有 error（本次未开启 strict，文件已写出）",
              file=sys.stderr)


def json_dump(obj):
    import json
    return json.dumps(obj, ensure_ascii=False, indent=1, default=str)


if __name__ == "__main__":
    main()
