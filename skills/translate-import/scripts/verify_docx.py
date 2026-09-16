# -*- coding: utf-8 -*-
"""Verify a generated DOCX: print stats JSON to stdout. Usage: python verify_docx.py <docx>"""
import json, sys
from docx import Document

def main():
    path = sys.argv[1]
    try:
        d = Document(path)
    except Exception as e:
        json.dump({"path": path, "ok": False, "error": str(e)[:200]}, sys.stdout)
        print(file=sys.stdout)
        return
    chars = 0
    paras = 0
    headings = 0
    for p in d.paragraphs:
        paras += 1
        chars += len(p.text.strip())
        if p.style.name.startswith("Heading"):
            headings += 1
    tables = len(d.tables)
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                chars += len(cell.text.strip())
    images = len(d.inline_shapes)
    json.dump({
        "path": path, "ok": chars > 500, "paras": paras, "chars": chars,
        "headings": headings, "tables": tables, "images": images,
    }, sys.stdout, ensure_ascii=False)
    print(file=sys.stdout)

if __name__ == "__main__":
    main()
