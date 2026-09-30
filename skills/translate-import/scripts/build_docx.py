# -*- coding: utf-8 -*-
"""Build a formatted academic translation DOCX from translation part files.

Usage:
    python build_docx.py <parts_dir> <out_docx>

<parts_dir> contains 01.json, 02.json, ... (sorted order). Each file:
    {"title_cn": "...",              # optional, first file only
     "blocks": [
        {"type":"title","text":"中文标题"},
        {"type":"subtitle","text":"English Original Title"},
        {"type":"heading","level":1,"text":"摘要"},
        {"type":"para","text":"..."},                       # body paragraph
        {"type":"para","text":"...","noindent":true},       # e.g. references entries
        {"type":"caption","text":"图1 ..."},
        {"type":"image","file":"figures/fig02_00.png","caption":"图1 ..."},   # file relative to parts_dir's parent
        {"type":"figure","file":"figures/fig02_00.png","caption":"图1 ..."},  # same as image, explicit figure
        {"type":"formula","latex":"$$E=mc^2$$","no":"1"},                    # LaTeX form
        {"type":"formula","latex":"$$E=mc^2$$","image":"formulas/fml02_00.png","no":"1"},
        {"type":"table","header":true,"rows":[["a","b"],["c","d"]]}
     ]}

图（figure / image）：优先原始图，其次 PDF 裁切图；两者都作为普通图片插入。
公式（formula / equation）：
  * latex  —— 数学公式的 LaTeX 写法，**转为 Word 原生公式对象（OMML）**
  * image  —— 原文公式截图，紧跟公式之后插入，供核对与保真
  * no     —— 原文式号，右对齐排在公式同一行

正文里的行内公式（`$E=mc^2$`）同样会被转成内联公式对象。

**LaTeX → OMML 依赖 pandoc**（见 omml.py）。**本流程不再静默降级**：
parts 里真的含公式、而 pandoc 不可用时，构建**直接报错退出**，不产出半成品——
这样「产出即终稿」才有保证。

**产物即终稿**：保存前会调用 `finalize.py` 做终稿规范化（图注规范、标题层级、
页眉 + 页脚 PAGE 域、正文引用换可点击跳转、参考文献一致性检查）。
任何一项不通过（尤其 pandoc 缺失、公式转换失败、参考文献对不上号）都会
**中止写出**，不会留下一份看起来成功、实际有问题的 DOCX。

Formatting: A4 single column; Chinese SimSun 10.5pt; Latin Times New Roman 10.5pt;
body justified, 1.5 line spacing, first-line indent 2 chars; headings bold black.
"""
import copy
import json, os, sys, glob
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from omml import LatexToOmml, clean_latex, segment_inline   # noqa: E402
from finalize import finalize_document                      # noqa: E402

EA_FONT = "宋体"
LATIN_FONT = "Times New Roman"

# 公式转换器：模块级单例，跨段共享缓存
_CONV = LatexToOmml()

def set_fonts(style, size_pt=None, bold=None, color_black=True):
    style.font.name = LATIN_FONT
    if size_pt is not None:
        style.font.size = Pt(size_pt)
    if bold is not None:
        style.font.bold = bold
    if color_black:
        style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), EA_FONT)

def run_fonts(run):
    run.font.name = LATIN_FONT
    r = run._element
    rpr = r.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), EA_FONT)

def setup_document(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(3.17)
    sec.right_margin = Cm(3.17)

    normal = doc.styles["Normal"]
    set_fonts(normal, 10.5)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)

    for name, size in (("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 10.5), ("Heading 4", 10.5)):
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        set_fonts(st, size, bold=True)
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing = 1.5
        st.paragraph_format.keep_with_next = True

def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run_fonts(run)
    run.bold = True
    run.font.size = Pt(15)

def add_subtitle(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(text)
    run_fonts(run)
    run.font.size = Pt(10.5)

def add_heading(doc, text, level):
    lvl = min(max(int(level), 1), 4)
    h = doc.add_heading("", level=lvl)
    run = h.add_run(text)
    run_fonts(run)

def _is_formula_para(text):
    """Display-formula paragraphs: LaTeX-heavy, little or no prose."""
    body = text.replace("$", " ").strip()
    cjk = sum(1 for ch in body if "\u4e00" <= ch <= "\u9fff")
    if cjk >= 2:
        return False
    if body.startswith("\\") or body.startswith("(") and "\\" in body:
        return True
    if "\\" in body and len(body) < 400:
        return True
    if text.count("$") >= 2 and len(body) < 120:
        return True
    return False

def add_para(doc, text, noindent=False):
    p = doc.add_paragraph()
    if _is_formula_para(text):
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if not noindent and not _is_formula_para(text):
        p.paragraph_format.first_line_indent = Pt(21)

    # 含行内公式：按 $...$ 切分，交替插入文本 run 与内联 OMML
    if "$" in text and _CONV.ok:
        segs = segment_inline(text)
        if any(k == "math" for k, _ in segs):
            _render_inline_math(p, segs)
            return p

    run = p.add_run(text)
    run_fonts(run)
    return p


def _render_inline_math(p, segs):
    """把 [("text", s)|("math", s)] 交替渲染进段落。

    行内公式转 OMML 成功则插入公式对象；失败则原样回退为 `$...$` 文本，
    保证内容不丢。

    公式对象与相邻汉字之间补一个空格，避免贴死；但**紧跟中文标点时不补**
    （否则会出现「公式 。」这种中文排版错误）。
    """
    # 中文标点：公式后面紧跟这些字符时不需要空格
    NO_SPACE_AFTER = "。，、；：？！）》」』】〉”’%,.;:!?)]}&"

    for i, (kind, val) in enumerate(segs):
        if kind == "text":
            if val:
                run_fonts(p.add_run(val))
            continue

        els = _CONV.convert(val, display=False)
        if not els:
            run_fonts(p.add_run("$%s$" % val))
            continue
        for el in els:
            p._element.append(copy.deepcopy(el))

        # 看下一个片段是否以中文标点开头
        nxt = segs[i + 1][1] if i + 1 < len(segs) else ""
        nxt = nxt.lstrip()
        if nxt and nxt[0] in NO_SPACE_AFTER:
            continue
        if not nxt:
            continue
        run_fonts(p.add_run(" "))

def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run_fonts(run)
    run.font.size = Pt(9)

def add_image(doc, path, caption):
    if not path or not os.path.exists(path):
        add_caption(doc, "（缺失图片：%s）" % (os.path.basename(path) if path else "未指定"))
        if caption:
            add_caption(doc, caption)
        return
    pic = doc.add_picture(path)
    width = pic.width if pic.width else 0
    height = pic.height if pic.height else 0
    max_w, max_h = Cm(14.6), Cm(20.0)
    if width > max_w or height > max_h:
        scale = min(max_w / width, max_h / height)
        pic.width = int(width * scale)
        pic.height = int(height * scale)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    if caption:
        add_caption(doc, caption)


def add_formula_para(doc, latex, image_path, no=None):
    """行间公式：Word 原生公式对象 + 原文截图 + 右对齐式号。

    两种产物并存：
      * latex —— 数学公式的 LaTeX 写法，转成 Word 原生公式对象（OMML）
                 转换失败时回退为纯文本 LaTeX，内容不丢
      * image —— 原文公式截图，紧跟其后插入，供逐字核对

    注意 feeds parts 里的 latex 常自带 `$$...$$` 定界符，需先剥掉再送 pandoc，
    否则会变成 `$$$$...$$$$` 导致解析失败。
    """
    if latex:
        # 剥定界符与式号（parts 里两者都可能写在 latex 字段内）
        pure, num_in_latex = clean_latex(latex)
        num = no or num_in_latex

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.keep_with_next = True

        els = _CONV.convert(pure, display=True)
        if els:
            for el in els:
                p._element.append(copy.deepcopy(el))
            if num:
                _add_formula_number(p, num)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            # 降级：纯文本 LaTeX + 式号（不丢内容）
            if num:
                _add_tabbed_text_formula(p, pure, num)
            else:
                run_fonts(p.add_run(pure))

    # 原文公式截图
    if image_path:
        if not os.path.exists(image_path):
            add_caption(doc, "（缺失公式截图：%s）" % os.path.basename(image_path))
        else:
            pic = doc.add_picture(image_path)
            width = pic.width if pic.width else 0
            height = pic.height if pic.height else 0
            max_w, max_h = Cm(13.0), Cm(6.0)
            if width > max_w or height > max_h:
                scale = min(max_w / width, max_h / height)
                pic.width = int(width * scale)
                pic.height = int(height * scale)
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.paragraphs[-1].paragraph_format.space_after = Pt(6)

    if not latex and not image_path:
        add_caption(doc, "（公式缺失：既无 LaTeX 也无截图）")


def _add_formula_number(p, no):
    """公式对象已在段中，右对齐补式号（制表位实现）。"""
    from docx.enum.text import WD_TAB_ALIGNMENT
    p.paragraph_format.tab_stops.add_tab_stop(Cm(14.6), WD_TAB_ALIGNMENT.RIGHT)
    run_fonts(p.add_run("\t(%s)" % no))


def _add_tabbed_text_formula(p, latex, no):
    """降级路径：左 LaTeX、右式号，用右对齐制表位实现。"""
    from docx.enum.text import WD_TAB_ALIGNMENT
    p.paragraph_format.tab_stops.add_tab_stop(Cm(14.6), WD_TAB_ALIGNMENT.RIGHT)
    run_fonts(p.add_run(latex))
    run_fonts(p.add_run("\t(%s)" % no))


def add_figure(doc, path, caption):
    """插图：优先原始图，其次 PDF 裁切图，二者统一走图片插入。"""
    add_image(doc, path, caption)

def add_table(doc, rows, header):
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncols)
    try:
        t.style = doc.styles["Table Grid"]
    except KeyError:
        pass
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    cell_size = Pt(9) if ncols <= 6 else Pt(8)
    total_w = Cm(14.6)
    col_w = int(total_w / ncols)
    for j in range(ncols):
        for i in range(len(rows)):
            t.cell(i, j).width = col_w
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = t.cell(i, j)
            txt = str(row[j]) if j < len(row) else ""
            p = cell.paragraphs[0]
            run = p.add_run(txt)
            run_fonts(run)
            run.font.size = cell_size
            if header and i == 0:
                run.bold = True
    doc.add_paragraph()

def _resolve(parts_dir, rel):
    """把块里的相对路径解析到 parts_dir 的父目录。"""
    if not rel:
        return ""
    if os.path.isabs(rel):
        return rel
    return os.path.join(os.path.dirname(parts_dir), rel)


def _has_formulas(parts_dir):
    """预扫描 parts：只有真存在公式时，才强制要求 pandoc 可用。

    没有公式的文档不该因为环境里没装 pandoc 就跑不动。
    """
    for fp in sorted(glob.glob(os.path.join(parts_dir, "*.json"))):
        try:
            with open(fp, encoding="utf-8") as f:
                data = json.load(f)
        except Exception:
            continue
        for b in data.get("blocks", []):
            if b.get("type") in ("formula", "equation"):
                return True
            t = b.get("text") or b.get("latex") or ""
            if "$" in t:
                return True
    return False


def build(parts_dir, out_path, finalize=True, header=None, xref=True,
          check_refs=True, allow_unused_refs=False, xref_style=False):
    doc = Document()
    setup_document(doc)

    # ---- pandoc 硬门槛：含公式却无法转 OMML 时直接失败，不产出降级件 ----
    if _has_formulas(parts_dir) and not _CONV.ok:
        raise SystemExit(
            "ERROR: parts 里含公式，但未找到 pandoc —— 无法把 LaTeX 转成 Word\n"
            "       原生公式对象。本流程要求「产出的译文 DOCX 直接就是 Word 原生\n"
            "       公式对象」，因此不再降级为纯文本 LaTeX，构建中止。\n"
            "       安装：conda install -c conda-forge pandoc\n"
            "             或 https://pandoc.org/installing.html")

    files = sorted(glob.glob(os.path.join(parts_dir, "*.json")))
    if not files:
        raise SystemExit("no part json files in " + parts_dir)
    unknown = set()
    doc_title = None
    for fp in files:
        with open(fp, encoding="utf-8") as f:
            data = json.load(f)
        blocks = data.get("blocks", [])
        for b in blocks:
            bt = b.get("type")
            if bt == "title":
                doc_title = doc_title or b["text"]
                add_title(doc, b["text"])
            elif bt == "subtitle":
                add_subtitle(doc, b["text"])
            elif bt == "heading":
                add_heading(doc, b["text"], b.get("level", 1))
            elif bt == "para":
                add_para(doc, b["text"], noindent=bool(b.get("noindent")))
            elif bt == "caption":
                add_caption(doc, b["text"])
            elif bt in ("image", "figure"):
                # 图：优先原始图，其次 PDF 裁切图；两者都走图片插入
                add_figure(doc, _resolve(parts_dir, b.get("file", "")),
                           b.get("caption", ""))
            elif bt in ("formula", "equation"):
                # 公式：Word 原生公式对象 + 原文截图，两种产物并存
                add_formula_para(
                    doc,
                    b.get("latex") or b.get("text") or "",
                    _resolve(parts_dir, b.get("image", "")),
                    no=b.get("no") or b.get("number"),
                )
            elif bt == "table":
                add_table(doc, b.get("rows", []), bool(b.get("header")))
            else:
                unknown.add(str(bt))
    if unknown:
        print("warning: unrendered block types:", sorted(unknown), file=sys.stderr)

    # ---- 公式转换小结：失败即中止，避免「看起来成功、其实是纯文本」----
    print("formula: " + _CONV.report(), file=sys.stderr)
    fails = _CONV.failures()
    if fails:
        print("以下公式未能转为 OMML：", file=sys.stderr)
        for lx, reason in fails[:20]:
            print("  - %r -> %s" % (lx[:70], reason), file=sys.stderr)
        if len(fails) > 20:
            print("  ... 另有 %d 条" % (len(fails) - 20), file=sys.stderr)
        raise SystemExit(
            "ERROR: %d 个公式未能转为 Word 原生公式对象，构建中止（未写出文件）。\n"
            "       请检查 parts 里的 LaTeX 写法后重跑。" % len(fails))

    # ---- 终稿规范化（保存前就地做，产出即终稿）----
    if finalize:
        stats = finalize_document(
            doc,
            title=(header or doc_title)[:80] if (header or doc_title) else None,
            strict=True,
            xref=xref,
            check_refs=check_refs,
            allow_unused_refs=allow_unused_refs,
            xref_style=xref_style,
        )
        rc = stats.get("refs_check") or {}
        print("finalize: marks=%s captions=%s headings=%s header=%s page=%s "
              "refs=%s cites=%s links=%s"
              % (stats.get("marks_removed"), stats.get("captions"),
                 stats.get("headings"), stats.get("header"),
                 stats.get("page_number"), rc.get("refs"), rc.get("cites"),
                 (stats.get("xref") or {}).get("links")),
              file=sys.stderr)
        for w in (rc.get("warnings") or []):
            print("warning: " + w, file=sys.stderr)

    doc.save(out_path)
    print("saved:", out_path, file=sys.stderr)


if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser(description="由 parts/*.json 构建终稿译文 DOCX")
    ap.add_argument("parts_dir")
    ap.add_argument("out_docx")
    ap.add_argument("--no-finalize", action="store_true",
                    help="只出内容稿，不做终稿规范化（不推荐）")
    ap.add_argument("--header", default=None, help="页眉标题（默认取首个 title 块）")
    ap.add_argument("--no-xref", action="store_true", help="不做正文引用交叉跳转")
    ap.add_argument("--no-ref-check", action="store_true",
                    help="跳过参考文献一致性检查（不推荐）")
    ap.add_argument("--allow-unused-refs", action="store_true",
                    help="文献列表中有未被正文引用的条目时不算错误")
    ap.add_argument("--hyperlink-style", action="store_true",
                    help="交叉引用套用蓝色下划线样式")
    a = ap.parse_args()
    build(a.parts_dir, a.out_docx,
          finalize=not a.no_finalize,
          header=a.header,
          xref=not a.no_xref,
          check_refs=not a.no_ref_check,
          allow_unused_refs=a.allow_unused_refs,
          xref_style=a.hyperlink_style)
