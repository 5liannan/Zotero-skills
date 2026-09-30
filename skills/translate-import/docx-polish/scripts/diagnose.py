# -*- coding: utf-8 -*-
"""诊断 DOCX：统计未转换 LaTeX / 高亮残留 / OMML 对象数 / 排版现状。

用法：
    python diagnose.py <file.docx>

同时用于「优化前诊断」与「优化后验收」——同一个脚本，看数字变化。
"""
import re
import sys

from docx import Document
from docx.oxml.ns import qn

M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"

DOLLAR_RE = re.compile(r"\$[^$\n]+\$")
BACKSLASH_CMD_RE = re.compile(r"\\(?:frac|sqrt|sum|int|alpha|beta|gamma|delta|"
                              r"lambda|sigma|omega|times|cdot|le|ge|neq|approx|"
                              r"partial|nabla|infty|left|right|begin|end|text)\b")


def diagnose(path):
    d = Document(path)
    paras = d.paragraphs

    hl_runs = 0
    for p in paras:
        for r in p.runs:
            rpr = r._element.find(qn("w:rPr"))
            if rpr is not None and rpr.find(qn("w:highlight")) is not None:
                hl_runs += 1

    omath = 0
    omath_para = 0
    for p in paras:
        n = len(p._element.findall(".//{%s}oMath" % M_NS))
        if n:
            omath += n
            omath_para += 1

    dollar_paras = [p.text for p in paras if DOLLAR_RE.search(p.text)]
    backslash_paras = [p.text for p in paras if BACKSLASH_CMD_RE.search(p.text)]

    with_pagenum = 0
    for sec in d.sections:
        for hp in sec.footer.paragraphs:
            if "PAGE" in hp._element.xml:
                with_pagenum += 1

    print("=" * 58)
    print("文件:", path)
    print("=" * 58)
    print("段落总数        :", len(paras))
    print("图片数          :", len(d.inline_shapes))
    print("-" * 58)
    print("高亮残留 run    :", hl_runs, "  (目标 0)")
    print("OMML 公式对象数 :", omath, "  跨", omath_para, "段  (目标 = 公式总数)")
    print("含 $ 定界符段落 :", len(dollar_paras), "  (目标 0)")
    print("含反斜杠命令段落:", len(backslash_paras), "  (目标 0)")
    print("页脚含 PAGE 域  :", with_pagenum, "/", len(d.sections))
    print("-" * 58)

    for label, items in (("残留 LaTeX($)", dollar_paras),
                         ("残留反斜杠命令", backslash_paras)):
        if items:
            print("\n[%s] 前 5 例：" % label)
            for t in items[:5]:
                print("   ", (t[:110] + "...") if len(t) > 110 else t)

    ok = (hl_runs == 0 and not dollar_paras and not backslash_paras and omath > 0)
    print("\n判定:", "OK" if ok else "未通过（见上方未达标的项）")
    return ok


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    sys.exit(0 if diagnose(sys.argv[1]) else 1)
