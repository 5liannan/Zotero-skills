# -*- coding: utf-8 -*-
"""从 DOCX 抽出行间/行内 LaTeX 公式，交给 pandoc 转 Word 原生 OMML。

用法：
    python extract_formulas.py <src.docx> <formulas.json> [--pandoc PATH]

产出 formulas.json：
    [{idx, para, kind, latex, num, omml, ok, err, orig}]
    kind="display"  整段就是一个公式
    kind="inline"   正文中夹杂 $...$，results 逐片段给 OMML
"""
import argparse
import json
import os
import re
import subprocess
import sys

from docx import Document

# 整段（去首尾空白后）就是 $...$ 或 $$...$$
DISPLAY_RE = re.compile(r"^\s*\${1,2}\s*(.+?)\s*\${1,2}\s*$", re.S)
# 正文里夹杂的 $...$
INLINE_RE = re.compile(r"\$([^$]+?)\$")
# 末尾式号 (2.1) / (12)
NUM_RE = re.compile(r"\((\d+(?:\.\d+)?)\)\s*$")
# 式号前的排版指令与标点
TAIL_CMD_RE = re.compile(r"(\\quad|\\qquad|\\;|\\,|\\!|,|\.|\s)+$")


def split_num(latex):
    """摘出末尾式号；顺带去掉残留的 \\quad / , / . 尾巴。

    **必须在送 pandoc 之前剥离**：否则式号会被排进 OMML，
    之后优化脚本再右对齐渲染一次，同一式号出现两遍——
    「(2.1) …… (2.1)」。
    """
    m = NUM_RE.search(latex)
    if not m:
        return latex, None
    pure = TAIL_CMD_RE.sub("", latex[:m.start()])
    return pure.strip(), m.group(1)


def latex_to_omml(latex, display, pandoc, workdir):
    """用 pandoc 把 LaTeX 转成 docx，再取回 <m:oMath> 所在段落的 XML。"""
    delim = "$$" if display else "$"
    tmp_md = os.path.join(workdir, "_f.md")
    tmp_docx = os.path.join(workdir, "_f.docx")
    try:
        with open(tmp_md, "w", encoding="utf-8") as f:
            f.write("%s%s%s\n" % (delim, latex, delim))
        p = subprocess.run([pandoc, tmp_md, "-o", tmp_docx],
                           capture_output=True, timeout=60)
        if p.returncode != 0:
            return False, p.stderr.decode("utf-8", "replace")[:200]
        for para in Document(tmp_docx).paragraphs:
            xml = para._element.xml
            if "oMath" in xml:
                return True, xml
        return False, "no oMath produced"
    except Exception as e:
        return False, "%s: %s" % (type(e).__name__, e)
    finally:
        for f in (tmp_md, tmp_docx):
            try:
                os.remove(f)
            except OSError:
                pass


def classify(text):
    t = text.strip()
    if not t or "$" not in t:
        return None
    m = DISPLAY_RE.match(t)
    if m:
        return "display", m.group(1).strip()
    return "inline", t


def main():
    ap = argparse.ArgumentParser(description="抽取 DOCX 中的 LaTeX 并转 OMML")
    ap.add_argument("src", help="输入 .docx")
    ap.add_argument("out", help="输出 formulas.json")
    ap.add_argument("--pandoc", default="pandoc",
                    help="pandoc 可执行文件路径（Windows 常随 Anaconda/conda 提供，用 where pandoc 定位）")
    args = ap.parse_args()

    workdir = os.path.dirname(os.path.abspath(args.out)) or "."
    doc = Document(args.src)

    records = []
    for pi, para in enumerate(doc.paragraphs):
        c = classify(para.text)
        if not c:
            continue
        kind, payload = c
        if kind == "display":
            pure, num = split_num(payload)
            records.append({"idx": len(records), "para": pi, "kind": "display",
                            "latex": pure, "num": num, "orig": para.text.strip()})
        else:
            records.append({"idx": len(records), "para": pi, "kind": "inline",
                            "latex_list": [m.group(1).strip()
                                           for m in INLINE_RE.finditer(payload)],
                            "orig": para.text.strip()})

    n_disp = sum(1 for r in records if r["kind"] == "display")
    n_inl = sum(1 for r in records if r["kind"] == "inline")
    n_frag = sum(len(r.get("latex_list", [])) for r in records if r["kind"] == "inline")
    total = n_disp + n_frag
    print("公式段:", len(records), " 行间:", n_disp,
          " 含行内公式段:", n_inl, "（片段", n_frag, "个）")
    print("公式总数:", total)

    # ---------- 批量转 OMML ----------
    ok = bad = 0
    for r in records:
        if r["kind"] == "display":
            good, res = latex_to_omml(r["latex"], True, args.pandoc, workdir)
            r["ok"] = good
            if good:
                r["omml"] = res
                ok += 1
            else:
                r["err"] = res
                bad += 1
                print("  [FAIL] para %d: %s" % (r["para"], res))
                print("         latex=%r" % r["latex"][:100])
        else:
            results = []
            for lx in r["latex_list"]:
                good, res = latex_to_omml(lx, False, args.pandoc, workdir)
                results.append({"latex": lx, "ok": good,
                                "omml": res if good else None,
                                "err": None if good else res})
                ok += 1 if good else 0
                if good:
                    pass
                else:
                    bad += 1
                    print("  [FAIL-inline] para %d: %s" % (r["para"], res))
                    print("         latex=%r" % lx[:100])
            r["results"] = results

    print("\n转换成功:", ok, " 失败:", bad, " / 共", total)
    with open(args.out, "w", encoding="utf-8") as f:
        json.dump(records, f, ensure_ascii=False, indent=1)
    print("写出:", args.out)
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
