# -*- coding: utf-8 -*-
"""终稿规范化的幂等性测试 —— 对同一份 DOCX 连续跑两次 finalize，第二次必须零变化。

跑法：
    python tests/test_idempotent.py

为什么需要它：`finalize` 会改写 run 结构、增删书签。非幂等的实现会：
  * 重复添加同名书签（`_RefList1` 出现两次 —— 重复书签名是无效 OOXML，Word 只认第一个）
  * 反复重建图注 run
  * 对已恢复成正文的段落再套一遍正文样式
这些都不会报错，只会悄悄把文档改坏，所以必须用「第二次零差异」来兜住。
"""
from __future__ import annotations

import json
import os
import shutil
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.join(os.path.dirname(HERE), "scripts")
sys.path.insert(0, SCRIPTS)

from docx import Document                      # noqa: E402
from docx.oxml.ns import qn                    # noqa: E402

M = "http://schemas.openxmlformats.org/officeDocument/2006/math"

FIXTURE = [
    {"type": "title", "text": "幂等性测试"},
    {"type": "heading", "level": 1, "text": "1 引言"},
    {"type": "para", "text": "正文引用文献[1]与[2, 3]。"},
    {"type": "heading", "level": 1, "text": "2 方法"},
    {"type": "formula", "latex": "$$E = mc^2,$$", "no": "1"},
    {"type": "para", "text": "其中 $m$ 为质量[2]。"},
    # 形似图注、实为正文 —— 会被恢复成正文样式（也要幂等）
    {"type": "caption", "text": "图 1 给出了系统的基本原理。"},
    {"type": "caption", "text": "图2 系统结构示意图。"},
    {"type": "heading", "level": 1, "text": "参考文献"},
    {"type": "para", "noindent": True,
     "text": "[1] Example, A.: Demo one, J. Demo, 1(1), 1-2, 2024."},
    {"type": "para", "noindent": True,
     "text": "[2] Sample, B.: Demo two, J. Demo, 1(2), 3-4, 2024."},
    {"type": "para", "noindent": True,
     "text": "[3] Placeholder, C.: Demo three, J. Demo, 1(3), 5-6, 2024."},
]


def full_text(p):
    return "".join(t.text or "" for t in p._p.iter(qn("w:t")))


def signature(path):
    d = Document(path)
    return {
        "paras": [full_text(p) for p in d.paragraphs],
        "align": [str(p.alignment) for p in d.paragraphs],
        "sizes": [[r.font.size.pt if r.font.size else None for r in p.runs]
                  for p in d.paragraphs],
        "styles": [p.style.name for p in d.paragraphs],
        "omath": sum(len(p._element.findall(".//{%s}oMath" % M))
                     for p in d.paragraphs),
        "links": sum(len(p._p.findall(qn("w:hyperlink"))) for p in d.paragraphs),
        "bookmarks": sorted(e.get(qn("w:name")) or ""
                            for e in d.element.body.iter(qn("w:bookmarkStart"))),
        "header": "".join(full_text(p)
                          for p in d.sections[0].header.paragraphs).strip(),
        "footer_page": any((it.text or "").upper().find("PAGE") >= 0
                           for it in d.sections[0].footer._element.iter(
                               qn("w:instrText"))),
    }


def main():
    import build_docx
    from finalize import finalize

    base = tempfile.mkdtemp(prefix="tf_idem_")
    try:
        parts = os.path.join(base, "parts")
        os.makedirs(parts)
        with open(os.path.join(parts, "01.json"), "w", encoding="utf-8") as f:
            json.dump({"title_cn": "幂等性测试", "blocks": FIXTURE},
                      f, ensure_ascii=False)
        docx = os.path.join(base, "out.docx")
        build_docx.build(parts, docx)          # 第一次规范化（构建内嵌）

        s1 = signature(docx)
        finalize(docx, backup=False)           # 第二次规范化
        s2 = signature(docx)
        finalize(docx, backup=False)           # 第三次规范化
        s3 = signature(docx)

        print("=" * 70)
        print("幂等性测试（临时目录 %s）" % base)
        print("=" * 70)
        print("段落 %d ；公式 %d ；超链接 %d ；书签 %d ；图注 %d"
              % (len(s3["paras"]), s3["omath"], s3["links"],
                 len(s3["bookmarks"]),
                 len([p for p in s3["paras"] if p.strip().startswith("图 ")])))
        dup = [x for x in s3["bookmarks"]
               if s3["bookmarks"].count(x) > 1 and x]
        if dup:
            print("!! 存在重复书签名: %s" % sorted(set(dup)))

        bad = 0
        for label, a, b in (("第2次 vs 第3次", s2, s3), ("第1次 vs 第3次", s1, s3)):
            print()
            print("--- %s ---" % label)
            for k in a:
                same = a[k] == b[k]
                if not same:
                    bad += 1
                note = ""
                if not same and isinstance(a[k], list):
                    for i, (x, y) in enumerate(zip(a[k], b[k])):
                        if x != y:
                            note = "  首个差异 @%d: %r -> %r" % (i, str(x)[:60],
                                                                str(y)[:60])
                            break
                    if len(a[k]) != len(b[k]):
                        note += "  长度 %d -> %d" % (len(a[k]), len(b[k]))
                elif not same:
                    note = "  %r -> %r" % (str(a[k])[:60], str(b[k])[:60])
                print("  %-12s %s%s" % (k, "OK" if same else "DIFF", note))

        # 第 2、3 次之间必须完全一致（这才是幂等性的定义）
        strictly_idempotent = all(s2[k] == s3[k] for k in s2)
        print()
        print("=" * 70)
        if strictly_idempotent and not dup:
            print("PASS：重复执行零变化，且无重复书签")
            return 0
        print("FAIL：非幂等或存在重复书签")
        return 1
    finally:
        shutil.rmtree(base, ignore_errors=True)


if __name__ == "__main__":
    sys.exit(main())
